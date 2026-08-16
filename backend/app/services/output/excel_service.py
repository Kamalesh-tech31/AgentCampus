"""
excel_service.py — Excel (.xlsx) output for Scribe using openpyxl.

Features:
  - Dynamic columns — driven by the actual keys in records, never hardcoded
  - Human-readable headings via heading_map.humanize()
  - Auto-sized columns
  - Styled header row (bold, background fill)
  - Separate "Analysis" sheet when metrics are present
  - Optional chart sheet embedding
  - Safe handling of non-serializable values (datetime, UUID, etc.)
"""

from __future__ import annotations

import io
import logging
import os
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

from app.services.output.heading_map import humanize

logger = logging.getLogger(__name__)

# ── Attempt imports — fail gracefully ──────────────────────────────────────────
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    _OPENPYXL_OK = True
except ImportError:
    _OPENPYXL_OK = False
    logger.error("[ExcelService] openpyxl not installed.")

# Output directory — created at runtime
_OUTPUT_DIR = Path(__file__).resolve().parents[3] / "output_files"


def _ensure_output_dir() -> Path:
    _OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    return _OUTPUT_DIR


def _safe_cell_value(value: Any) -> Any:
    """Convert a value to something openpyxl can write safely."""
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, datetime):
        return str(value)
    # Dicts / lists (e.g. nested department_breakdown) → JSON-like string
    import json
    try:
        return json.dumps(value, ensure_ascii=False)
    except Exception:
        return str(value)


def _apply_header_style(cell, fill_color: str = "1F4E79") -> None:
    """Apply bold white text on dark blue background to a header cell."""
    cell.font = Font(bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)


def _apply_border(cell) -> None:
    thin = Side(border_style="thin", color="D3D3D3")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _write_records_sheet(ws, records: list[dict]) -> None:
    """Write student records to the given worksheet with dynamic columns."""
    if not records:
        ws["A1"] = "No records available."
        return

    # Determine column order from first record, exclude only the raw 'id' field
    keys = [k for k in records[0].keys() if k not in {"id"}]
    headers = [humanize(k) for k in keys]

    # ── Header row ────────────────────────────────────────────────────────────
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        _apply_header_style(cell)
        _apply_border(cell)

    ws.row_dimensions[1].height = 20

    # ── Data rows ─────────────────────────────────────────────────────────────
    alt_fill = PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid")
    for row_idx, record in enumerate(records, start=2):
        fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, key in enumerate(keys, start=1):
            cell = ws.cell(
                row=row_idx,
                column=col_idx,
                value=_safe_cell_value(record.get(key)),
            )
            cell.alignment = Alignment(horizontal="left", vertical="center")
            if fill:
                cell.fill = fill
            _apply_border(cell)

    # ── Auto-size columns ─────────────────────────────────────────────────────
    for col_idx in range(1, len(keys) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            len(str(ws.cell(row=r, column=col_idx).value or ""))
            for r in range(1, len(records) + 2)
        )
        ws.column_dimensions[col_letter].width = min(max_len + 3, 35)

    # Freeze header row
    ws.freeze_panes = "A2"


def _write_analysis_sheet(ws, metrics: dict, insight: Optional[str]) -> None:
    """Write metrics summary to the Analysis worksheet."""

    def _kv_row(row: int, label: str, value: Any, label_fill: str = "2E4057") -> None:
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = Font(bold=True, color="FFFFFF", size=10)
        lc.fill = PatternFill(start_color=label_fill, end_color=label_fill, fill_type="solid")
        lc.alignment = Alignment(horizontal="left", vertical="center")
        _apply_border(lc)

        vc = ws.cell(row=row, column=2, value=_safe_cell_value(value))
        vc.alignment = Alignment(horizontal="left", vertical="center")
        _apply_border(vc)
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 25

    # ── Title ─────────────────────────────────────────────────────────────────
    ws.merge_cells("A1:B1")
    title_cell = ws["A1"]
    title_cell.value = "Analytics Summary"
    title_cell.font = Font(bold=True, size=13, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    row = 3
    scalar_fields = [
        ("totalRecords", "total_records", "Total Records"),
        ("averageCgpa", "average_cgpa", "Average CGPA"),
        ("highestCgpa", "highest_cgpa", "Highest CGPA"),
        ("lowestCgpa", "lowest_cgpa", "Lowest CGPA"),
        ("avgAttendance", "avg_attendance", "Avg Attendance (%)"),
        ("probationCount", "probation_count", "Students on Probation"),
    ]
    for camel, snake, label in scalar_fields:
        value = metrics.get(camel) or metrics.get(snake)
        if value is not None:
            val_str = f"{value:.2f}" if isinstance(value, float) else str(value)
            _kv_row(row, label, val_str)
            row += 1

    if insight:
        row += 1
        ws.merge_cells(f"A{row}:B{row}")
        insight_cell = ws[f"A{row}"]
        insight_cell.value = f"Insight: {insight}"
        insight_cell.font = Font(italic=True, size=10, color="333333")
        insight_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 36
        row += 1

    # Department breakdown sub-table
    breakdown = metrics.get("departmentBreakdown") or metrics.get("department_breakdown")
    if breakdown and isinstance(breakdown, dict):
        row += 1
        ws.merge_cells(f"A{row}:B{row}")
        hc = ws[f"A{row}"]
        hc.value = "Department Breakdown"
        hc.font = Font(bold=True, size=11, color="FFFFFF")
        hc.fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
        hc.alignment = Alignment(horizontal="center")
        row += 1

        # Sub-header
        ws.cell(row=row, column=1, value="Department").font = Font(bold=True)
        ws.cell(row=row, column=2, value="Avg CGPA").font = Font(bold=True)
        ws.cell(row=row, column=3, value="Count").font = Font(bold=True)
        ws.column_dimensions["C"].width = 12
        row += 1

        for dept, data in breakdown.items():
            if isinstance(data, dict):
                avg = data.get("avgCgpa") or data.get("avg_cgpa", 0.0)
                cnt = data.get("count", 0)
                ws.cell(row=row, column=1, value=dept)
                ws.cell(row=row, column=2, value=round(float(avg), 2) if avg else 0.0)
                ws.cell(row=row, column=3, value=cnt)
                row += 1


def generate_excel(
    records: list[dict],
    metrics: Optional[dict],
    insight: Optional[str],
    file_stem: str = "student_report",
    pulse_data: Optional[dict] = None,
) -> dict:
    """
    Generate an Excel file and return file metadata.

    Args:
        records:    List of student record dicts from Vault.
        metrics:    OrchestrationMetrics dict from Pulse (or None).
        insight:    Natural-language insight string (or None).
        file_stem:  Base name for the output file (without extension).

    Returns:
        dict with keys: file_path, file_name, sheet_count

    Raises:
        RuntimeError: If openpyxl is not installed.
    """
    if not _OPENPYXL_OK:
        raise RuntimeError("openpyxl is required for Excel generation but is not installed.")

    wb = openpyxl.Workbook()
    sheet_count = 0

    # ── Specialized Pulse Data Multi-Sheet Workbook ───────────────────────────
    if pulse_data:
        # 1. Custom primary tables from Pulse
        if pulse_data.get("at_risk") or pulse_data.get("tables", {}).get("at_risk_students"):
            at_risk_recs = pulse_data.get("at_risk") or pulse_data.get("tables", {}).get("at_risk_students", [])
            ws_risk = wb.active
            ws_risk.title = "At-Risk Students"
            _write_records_sheet(ws_risk, at_risk_recs)
            sheet_count += 1

            # Department breakdown sheet if present
            dept_recs = pulse_data.get("department_analysis") or pulse_data.get("tables", {}).get("department_risk", [])
            if dept_recs:
                ws_dept = wb.create_sheet(title="Department Risk")
                _write_records_sheet(ws_dept, dept_recs)
                sheet_count += 1

        elif pulse_data.get("ranking") or pulse_data.get("tables", {}).get("top_performers"):
            top_recs = pulse_data.get("ranking") or pulse_data.get("tables", {}).get("top_performers", [])
            ws_rank = wb.active
            ws_rank.title = "Top Performers"
            _write_records_sheet(ws_rank, top_recs)
            sheet_count += 1

        elif pulse_data.get("tables"):
            active_set = False
            for tname, trows in pulse_data["tables"].items():
                if trows:
                    sheet_title = humanize(tname)[:31]
                    if not active_set:
                        ws_tbl = wb.active
                        ws_tbl.title = sheet_title
                        active_set = True
                    else:
                        ws_tbl = wb.create_sheet(title=sheet_title)
                    _write_records_sheet(ws_tbl, trows)
                    sheet_count += 1
        else:
            ws_data = wb.active
            ws_data.title = "Student Data"
            _write_records_sheet(ws_data, records)
            sheet_count += 1

        # Embed dynamic charts from Pulse
        if pulse_data.get("chart_data"):
            from app.services.output.chart_service import render_dynamic_chart
            from openpyxl.drawing.image import Image as XLImage
            ws_chart = wb.create_sheet(title="Visualizations")
            row_anchor = 1
            for cname, cdata in pulse_data["chart_data"].items():
                try:
                    cbytes = render_dynamic_chart(cdata)
                    if cbytes:
                        img = XLImage(io.BytesIO(cbytes))
                        img.anchor = f"A{row_anchor}"
                        ws_chart.add_image(img)
                        row_anchor += 25
                        sheet_count += 1
                except Exception as exc:
                    logger.warning(f"[ExcelService] Could not embed dynamic chart {cname}: {exc}")

    else:
        # ── Legacy Sheet 1: Student Data ──────────────────────────────────────
        ws_data = wb.active
        ws_data.title = "Student Data"
        _write_records_sheet(ws_data, records)
        sheet_count += 1

        # ── Legacy Sheet 2: Analysis (optional) ───────────────────────────────
        if metrics:
            ws_analysis = wb.create_sheet(title="Analysis")
            _write_analysis_sheet(ws_analysis, metrics, insight)
            sheet_count += 1

        # ── Legacy Chart Sheet ────────────────────────────────────────────────
        if records:
            try:
                from app.services.output.chart_service import cgpa_distribution_chart
                from openpyxl.drawing.image import Image as XLImage

                chart_bytes = cgpa_distribution_chart(records)
                if chart_bytes:
                    ws_chart = wb.create_sheet(title="CGPA Chart")
                    img = XLImage(io.BytesIO(chart_bytes))
                    img.anchor = "A1"
                    ws_chart.add_image(img)
                    sheet_count += 1
            except Exception as exc:
                logger.warning(f"[ExcelService] Could not embed chart: {exc}")

    # ── Save ──────────────────────────────────────────────────────────────────
    out_dir = _ensure_output_dir()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_name = f"{file_stem}_{timestamp}.xlsx"
    file_path = str(out_dir / file_name)

    wb.save(file_path)
    logger.info(f"[ExcelService] Saved Excel report: {file_path}")

    return {
        "file_path": file_path,
        "file_name": file_name,
        "sheet_count": sheet_count,
    }
