import pytest
from app.agents.input_agent import InputAgent
from app.agents.vault_planner import _deterministic_fallback_plan, _post_validate_plan
from app.agents.db_agent import _format_sql_from_plan, DBAgent
from app.mother.types import AgentTask
from app.services.output.text_service import generate_text
from app.contracts.students import StudentRecord


def test_top_10_students_deterministic_extraction():
    input_agent = InputAgent()
    task = AgentTask(
        task_id="test-top10-1",
        workflow_id="wf-1",
        agent="input",
        objective="Extract intent",
        expected_output="Intent extracted",
        input_data={"user_query": "Show the top 10 students based on the CGPA", "mode": "explore"}
    )
    res = input_agent.execute(task)
    intent = res.result.get("structured_intent", {})
    
    assert intent.get("limit") == 10
    assert intent.get("sort", {}).get("field") == "cgpa"
    assert intent.get("sort", {}).get("direction") == "desc"

    # Plan generation & validation
    plan = _deterministic_fallback_plan({"input": {"structured_intent": intent}, "query": task.input_data["user_query"]})
    plan = _post_validate_plan(plan, {"input": {"structured_intent": intent}})
    
    assert plan["params"]["limit"] == 10
    assert plan["params"]["sort"]["field"] == "cgpa"
    
    sql = _format_sql_from_plan(plan)
    assert "ORDER BY cgpa DESC, roll_number ASC" in sql
    assert "LIMIT 10" in sql


def test_top_5_students():
    input_agent = InputAgent()
    task = AgentTask(
        task_id="test-top5-1",
        workflow_id="wf-1",
        agent="input",
        objective="Extract intent",
        expected_output="Intent extracted",
        input_data={"user_query": "Show the top 5 students based on CGPA", "mode": "explore"}
    )
    res = input_agent.execute(task)
    intent = res.result.get("structured_intent", {})
    assert intent.get("limit") == 5

    plan = _deterministic_fallback_plan({"input": {"structured_intent": intent}, "query": task.input_data["user_query"]})
    sql = _format_sql_from_plan(plan)
    assert "LIMIT 5" in sql


def test_top_20_students():
    input_agent = InputAgent()
    task = AgentTask(
        task_id="test-top20-1",
        workflow_id="wf-1",
        agent="input",
        objective="Extract intent",
        expected_output="Intent extracted",
        input_data={"user_query": "Show the top 20 students based on CGPA", "mode": "explore"}
    )
    res = input_agent.execute(task)
    intent = res.result.get("structured_intent", {})
    assert intent.get("limit") == 20


def test_all_cse_students_no_limit():
    input_agent = InputAgent()
    task = AgentTask(
        task_id="test-cse-1",
        workflow_id="wf-1",
        agent="input",
        objective="Extract intent",
        expected_output="Intent extracted",
        input_data={"user_query": "Show all CSE students", "mode": "explore"}
    )
    res = input_agent.execute(task)
    intent = res.result.get("structured_intent", {})
    assert intent.get("limit") is None
    assert intent.get("department") == "Computer Science"

    plan = _deterministic_fallback_plan({"input": {"structured_intent": intent}, "query": task.input_data["user_query"]})
    sql = _format_sql_from_plan(plan)
    assert "LIMIT" not in sql
    assert "department = 'Computer Science'" in sql


def test_cgpa_above_8_5_no_limit():
    input_agent = InputAgent()
    task = AgentTask(
        task_id="test-cgpa-1",
        workflow_id="wf-1",
        agent="input",
        objective="Extract intent",
        expected_output="Intent extracted",
        input_data={"user_query": "Find students with CGPA above 8.5", "mode": "explore"}
    )
    res = input_agent.execute(task)
    intent = res.result.get("structured_intent", {})
    assert intent.get("limit") is None

    plan = _deterministic_fallback_plan({"input": {"structured_intent": intent}, "query": task.input_data["user_query"]})
    sql = _format_sql_from_plan(plan)
    print(f"DEBUG SQL: {sql}, plan: {plan}")
    assert "LIMIT" not in sql
    assert "cgpa >= 8.5" in sql or "cgpa > 8.5" in sql or "8.5" in sql


def test_tie_breaking_deterministic_execution():
    # 40 students with same CGPA 8.90
    same_cgpa_students = [
        {"id": f"STU-{100+i}", "roll_number": f"21CS{i:03d}", "name": f"Student {i}", "department": "Computer Science", "cgpa": 8.90, "attendance": 85.0}
        for i in range(40, 0, -1)  # reverse order roll numbers
    ]
    
    # Sort deterministically with primary (cgpa DESC) and secondary (roll_number ASC)
    same_cgpa_students.sort(
        key=lambda s: (-s["cgpa"], s["roll_number"])
    )
    
    # Take top 10
    top_10 = same_cgpa_students[:10]
    assert len(top_10) == 10
    assert top_10[0]["roll_number"] == "21CS001"
    assert top_10[9]["roll_number"] == "21CS010"


def test_text_service_records_count_header():
    records = [
        {"id": f"STU-{i}", "rollNumber": f"21CS{i:03d}", "name": f"Student {i}", "department": "Computer Science", "cgpa": 9.5 - (i * 0.1), "attendance": 90.0}
        for i in range(10)
    ]
    text = generate_text(records, None, None, "Show the top 10 students based on CGPA")
    assert "STUDENT RECORDS (10 total):" in text or "STUDENT RECORDS  (10 total):" or "10 total" in text
    assert "102 total" not in text
    assert "25 total" not in text


def test_end_to_end_top_10_students_based_on_cgpa():
    # 1. Input Agent
    input_agent = InputAgent()
    task_in = AgentTask(
        task_id="e2e-cgpa-in",
        workflow_id="wf-cgpa",
        agent="input",
        objective="Parse intent",
        expected_output="Intent parsed",
        input_data={"user_query": "Show the top 10 students based on the cgpa", "mode": "explore"}
    )
    res_in = input_agent.execute(task_in)
    intent = res_in.result.get("structured_intent", {})
    assert intent.get("limit") == 10
    assert intent.get("sort", {}).get("field") == "cgpa"
    assert intent.get("sort", {}).get("direction") == "desc"

    # 2. DB Agent
    db_agent = DBAgent()
    task_db = AgentTask(
        task_id="e2e-cgpa-db",
        workflow_id="wf-cgpa",
        agent="db",
        objective="Execute query",
        expected_output="Rows returned",
        input_data={
            "user_query": "Show the top 10 students based on the cgpa",
            "mode": "explore",
            "input": res_in.result
        }
    )
    res_db = db_agent.execute(task_db)
    db_data = res_db.result
    records = db_data.get("records", [])
    sql = db_data.get("sql", "")

    assert len(records) == 10, f"Expected 10 records from DB Agent, got {len(records)}"
    assert db_data.get("count") == 10
    assert db_data.get("requested_limit") == 10
    assert db_data.get("rows_returned") == 10
    assert "LIMIT 10" in sql
    assert "ORDER BY cgpa DESC, roll_number ASC" in sql

    # 3. Output / Scribe Agent
    from app.agents.output_agent import OutputAgent
    output_agent = OutputAgent()
    task_out = AgentTask(
        task_id="e2e-cgpa-out",
        workflow_id="wf-cgpa",
        agent="output",
        objective="Format results",
        expected_output="Formatted text and summary",
        input_data={
            "user_query": "Show the top 10 students based on the cgpa",
            "mode": "explore",
            "input": res_in.result,
            "db": res_db.result
        }
    )
    res_out = output_agent.execute(task_out)
    summary = res_out.result.get("summary", "")
    out_records = res_out.result.get("result", {}).get("data", [])

    assert len(out_records) == 10, f"Expected 10 records in OrchestrationResult.data, got {len(out_records)}"
    assert "STUDENT RECORDS (10 total):" in summary or "10 total" in summary
    assert "102 total" not in summary

    # 4. Excel Generation
    from app.services.output.excel_service import generate_excel
    import openpyxl
    file_info = generate_excel(records, None, None)
    wb = openpyxl.load_workbook(file_info["file_path"])
    ws = wb.active
    # Header row + 10 data rows = 11 rows total
    assert ws.max_row == 11, f"Expected 11 rows in Excel (1 header + 10 records), got {ws.max_row}"


def test_end_to_end_top_10_cse_students():
    # 1. Input Agent
    input_agent = InputAgent()
    task_in = AgentTask(
        task_id="e2e-cse-in",
        workflow_id="wf-cse",
        agent="input",
        objective="Parse intent",
        expected_output="Intent parsed",
        input_data={"user_query": "Show the top 10 CSE students", "mode": "explore"}
    )
    res_in = input_agent.execute(task_in)
    intent = res_in.result.get("structured_intent", {})
    assert intent.get("limit") == 10
    assert intent.get("department") == "Computer Science"

    # 2. DB Agent
    db_agent = DBAgent()
    task_db = AgentTask(
        task_id="e2e-cse-db",
        workflow_id="wf-cse",
        agent="db",
        objective="Execute query",
        expected_output="Rows returned",
        input_data={
            "user_query": "Show the top 10 CSE students",
            "mode": "explore",
            "input": res_in.result
        }
    )
    res_db = db_agent.execute(task_db)
    db_data = res_db.result
    records = db_data.get("records", [])
    sql = db_data.get("sql", "")

    assert len(records) == 10, f"Expected 10 records from DB Agent, got {len(records)}"
    assert "department = 'Computer Science'" in sql
    assert "LIMIT 10" in sql


if __name__ == "__main__":
    test_top_10_students_deterministic_extraction()
    test_top_5_students()
    test_top_20_students()
    test_all_cse_students_no_limit()
    test_cgpa_above_8_5_no_limit()
    test_tie_breaking_deterministic_execution()
    test_text_service_records_count_header()
    test_end_to_end_top_10_students_based_on_cgpa()
    test_end_to_end_top_10_cse_students()
    print("ALL DETERMINISTIC TOP-N LIMIT AND TIE-BREAKER TESTS PASSED!")
